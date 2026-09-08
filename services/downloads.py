"""Contained, read-only downloads. Local operators remain trusted (no v0 auth)."""
from io import BytesIO
import os
import re
import stat
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from pptx import Presentation

from services.artifacts import OUTPUT
from services.proof import ProofManifest
from services.query import PermissionContext

MAX_BYTES = 16 * 1024 * 1024
UUID_PATTERN = r'[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}'


def read_artifact(name: str, permission: PermissionContext) -> bytes:
    """Never follow links; check the exact bytes delivered, not a reopened path."""
    if not re.fullmatch(rf'{UUID_PATTERN}(?:-{UUID_PATTERN})?\.(md|xlsx|pptx)', name):
        raise ValueError('Use an artifact basename from output-workspace.')
    directory = os.open(OUTPUT, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        fd = os.open(name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
        with os.fdopen(fd, 'rb') as stream:
            info = os.fstat(stream.fileno())
            if not stat.S_ISREG(info.st_mode) or info.st_size > MAX_BYTES:
                raise ValueError('Artifact is not a bounded regular file.')
            payload = stream.read(MAX_BYTES + 1)
    finally:
        os.close(directory)
    if len(payload) > MAX_BYTES:
        raise ValueError('Artifact exceeds download size limit.')
    try:
        proofs: list[ProofManifest]
        if name.endswith('.md'):
            text = payload.decode('utf-8')
            proof_text = text.rsplit('\n## Proof manifest\n```json\n', 1)[1]
            proofs = [ProofManifest.model_validate_json(proof_text.removesuffix('```\n').rstrip())]
        else:
            with ZipFile(BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_BYTES:
                    raise ValueError('Expanded artifact exceeds download size limit.')
            if name.endswith('.xlsx'):
                book = load_workbook(BytesIO(payload), read_only=True, data_only=False)
                try:
                    proof_text = ''.join(str(row[0]) for row in book['Manifest'].values)
                    proofs = [ProofManifest.model_validate_json(proof_text)]
                finally:
                    book.close()
            else:
                deck = Presentation(BytesIO(payload))
                proofs = []
                for slide in deck.slides:
                    notes = slide.notes_slide.notes_text_frame
                    if notes is None:
                        raise ValueError('Missing artifact proof.')
                    proofs.append(ProofManifest.model_validate_json(notes.text))
        if not proofs or any(p.permission_context != permission.model_dump() for p in proofs):
            raise ValueError('Artifact permission does not match the caller.')
        if any(p != proofs[0] for p in proofs):
            raise ValueError('Artifact proof is inconsistent.')
    except (KeyError, IndexError, UnicodeError, BadZipFile) as exc:
        raise ValueError('Artifact has no readable proof manifest.') from exc
    return payload
