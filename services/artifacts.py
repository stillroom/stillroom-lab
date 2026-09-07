"""Local-only artifact boundary. No source access, model calls, or arithmetic."""
from pathlib import Path
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
import os
import re
import textwrap

from pptx import Presentation
from pptx.util import Inches, Pt

from services.query import QueryResult

OUTPUT = Path(__file__).resolve().parents[1] / 'output-workspace'


def literal(text: str) -> str:
    """Keep retrieved text as data in Markdown, including HTML and footnote syntax."""
    return re.sub(r'([\\`*_{}\[\]<>()#!|])', r'\\\1', text).replace('\n', ' ')


def briefing(result: QueryResult) -> str:
    proof = result.manifest
    refs = ' '.join(f'[^{literal(s)}]' for s in proof.source_ids)
    lines = [f'# {proof.approved_query}', f'As of: {proof.as_of}; query: {result.query_id}',
             f'Snapshot bundle: {proof.snapshot_id}', result.reason.message,
             'Completeness check: not_established. Observations are not complete totals.']
    for row in result.rows:
        lines.append(f'- {literal(row.customer_id)}: observed {row.observed_amount}; '
                     f'confirmed amount: {row.amount if row.amount is not None else "not established"}. {refs}')
    for source in proof.sources:
        if source.ticket is not None:
            lines.append(f'- Unresolved ticket {literal(source.canonical_id)}: '
                         f'{literal(source.ticket.status)} — {literal(source.ticket.subject)} '
                         f'[^{literal(source.source_id)}]')
    lines.extend(f'- Caveat: {literal(c)}' for c in proof.caveats)
    lines.extend(f'- Conflict: {literal(c)}' for c in proof.unresolved_conflicts)
    for source in proof.sources:
        lines.append(f'[^{literal(source.source_id)}]: {source.entity} {literal(source.canonical_id)}; '
                     f'snapshot {source.snapshot_id}; observed {source.observed_at}; '
                     f'retrieved {source.retrieved_at.isoformat()}.')
    lines.extend(['', '## Proof manifest', '```json', proof.model_dump_json(indent=2), '```'])
    return '\n'.join(lines) + '\n'


def text_row(sheet: Worksheet, values: list[str]) -> None:
    # Force strings: source text must never become a formula, and Decimal must
    # never enter openpyxl's lossy numeric serializer. Fail instead of truncating.
    if any(len(value) > 32767 for value in values):
        raise ValueError('Evidence exceeds spreadsheet cell capacity')
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.data_type = 's'


def spreadsheet(result: QueryResult) -> bytes:
    book = Workbook()
    book.remove(book.worksheets[0])
    rows = book.create_sheet('Results')
    text_row(rows, ['Customer', 'Client', 'Observed amount (exact decimal text)', 'Confirmed amount'])
    for row in result.rows:
        text_row(rows, [row.customer_id, row.client_id, str(row.observed_amount),
                        str(row.amount) if row.amount is not None else 'not established'])
    tickets = book.create_sheet('Tickets')
    text_row(tickets, ['Ticket', 'Status', 'Subject', 'Source ID'])
    for source in result.manifest.sources:
        if source.ticket is not None:
            text_row(tickets, [source.canonical_id, source.ticket.status,
                              source.ticket.subject, source.source_id])
    sources = book.create_sheet('Sources')
    text_row(sources, ['Source ID', 'Snapshot ID', 'Retrieval time', 'Observed date', 'Caveats'])
    for source in result.manifest.sources:
        text_row(sources, [source.source_id, str(source.snapshot_id), source.retrieved_at.isoformat(),
                           str(source.observed_at), '; '.join(result.manifest.caveats + source.conflicts)])
    proof = book.create_sheet('Manifest')
    payload = result.manifest.model_dump_json()
    for start in range(0, len(payload), 30000):
        text_row(proof, [payload[start:start + 30000]])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def presentation(result: QueryResult) -> bytes:
    deck = Presentation()
    deck.slide_width = Inches(12)
    deck.slide_height = Inches(7)
    proof = result.manifest
    messages = [*proof.caveats, 'Completeness check: not_established.']
    messages.extend(f'{row.customer_id}: observed {row.observed_amount}; confirmed amount: '
                    f'{row.amount if row.amount is not None else "not established"}' for row in result.rows)
    messages.extend(f'Unresolved ticket {s.canonical_id}: {s.ticket.status} — {s.ticket.subject}'
                    for s in proof.sources if s.ticket is not None)
    messages.extend(proof.unresolved_conflicts)
    for message in messages:
        lines = textwrap.wrap(message, width=80) or ['']
        for start in range(0, len(lines), 8):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(11), Inches(1)).text_frame
            title.text = str(proof.approved_query)
            title.paragraphs[0].font.size = Pt(26)
            body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(11), Inches(4)).text_frame
            body.text = '\n'.join(lines[start:start + 8])
            for paragraph in body.paragraphs:
                paragraph.font.size = Pt(22)
            footer = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(11), Inches(0.7)).text_frame
            footer.text = f'As of {proof.as_of} | Snapshot {proof.snapshot_id} | Evidence in notes'
            footer.paragraphs[0].font.size = Pt(12)
            notes = slide.notes_slide.notes_text_frame
            if notes is None:
                raise ValueError('Template has no evidence notes')
            notes.text = proof.model_dump_json()
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def generate(result: QueryResult, target: Path) -> Path:
    """Accept validated governed results; never overwrite existing paths."""
    result = QueryResult.model_validate(result.model_dump())
    proof = result.manifest
    if proof.snapshot_id != result.snapshot_id:
        raise ValueError('Result and proof snapshot differ')
    if set(proof.source_ids) != {s.source_id for s in proof.sources}:
        raise ValueError('Proof source coverage is inconsistent')
    if not {c for s in proof.sources for c in s.conflicts} <= set(proof.unresolved_conflicts):
        raise ValueError('Unresolved conflicts cannot be dropped')
    for row in result.rows:
        if (proof.permission_context is None or row.client_id != proof.permission_context.get('client_id')
                or row.customer_id not in {s.canonical_id for s in proof.sources if s.entity == 'Customer'}):
            raise ValueError('Rows do not match the evidence permission context')
    target = target.absolute()
    if '..' in target.parts or not target.is_relative_to(OUTPUT):
        raise ValueError('Artifacts must stay inside output-workspace')
    if target.suffix == '.md':
        payload = briefing(result).encode('utf-8')
    elif target.suffix == '.xlsx':
        payload = spreadsheet(result)
    elif target.suffix == '.pptx':
        payload = presentation(result)
    else:
        raise ValueError('Unsupported artifact extension')
    OUTPUT.mkdir(exist_ok=True)
    directory = os.open(OUTPUT, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        for part in target.relative_to(OUTPUT).parts[:-1]:
            try:
                os.mkdir(part, dir_fd=directory)
            except FileExistsError:
                pass
            child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=directory)
            os.close(directory)
            directory = child
        fd = os.open(target.name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                     0o600, dir_fd=directory)
        with os.fdopen(fd, 'wb') as stream:
            stream.write(payload)
    finally:
        os.close(directory)
    return target
