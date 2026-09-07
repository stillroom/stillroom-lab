"""Ticket 04 acceptance at query and artifact boundaries."""
from datetime import datetime, timezone
from pathlib import Path
import json
from datetime import date
import pytest

from services.query import QueryResult, QueryRequest, QueryService, PermissionContext

from fastapi.testclient import TestClient

from services.api import create_app
from services.store import Store
from test_ingestion import FAMILIES, approved_app


@pytest.mark.parametrize('scenario', ['revenue', 'pipeline', 'identity', 'failure'])
def test_empty_evidence_preserves_selected_mapping_versions(store: Store, scenario: str) -> None:
    from test_query_isolation import add_observation

    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    if scenario == 'identity':
        add_observation(store, 'Payment', 'P001', 'PX', None, {'invoice_id': 'OTHER-CLIENT-INVOICE'})
    with TestClient(create_app(app)) as client:
        if scenario == 'failure':
            with store.transaction() as db:
                db.execute('REVOKE SELECT ON canonical_records FROM lab_query')
        metric = scenario if scenario in ('revenue', 'pipeline') else 'overdue_balances'
        response = client.post('/queries', json={
            'approved_query': f'v0.1.0/{metric}',
            'permission': {'role': 'owner', 'client_id': 'CLIENT-1'}, 'as_of': '2026-09-01'})
    assert response.status_code == (503 if scenario == 'failure' else 200)
    body = response.json()
    assert body['manifest']['sources'] == []
    with store.transaction() as db:
        audit = db.execute('SELECT * FROM query_audit WHERE query_id=%s', (body['query_id'],)).fetchone()
    assert audit is not None and audit['retrieval_started'] is True
    expected = ['billing:v0.1.0', 'crm:v0.1.0', 'email:v0.1.0', 'policy:v0.1.0', 'support:v0.1.0']
    assert audit['mapping_versions'] == expected
    assert body['manifest']['mapping_versions'] == expected


@pytest.mark.parametrize('permission', [
    {'role': 'owner', 'client_id': 'CLIENT-1'}, None, {},
    {'role': 42, 'client_id': 'CLIENT-1'},
    {'role': 'owner', 'client_id': 'CLIENT-1', 'extra': 'invalid'},
])
def test_malformed_date_preserves_only_valid_permission(store: Store, permission: object) -> None:
    with TestClient(create_app(approved_app(store))) as client:
        response = client.post('/queries', json={
            'approved_query': 'v0.1.0/cash_received', 'permission': permission, 'as_of': 'invalid'})
    assert response.status_code == 422
    body = response.json()
    expected = permission if permission == {'role': 'owner', 'client_id': 'CLIENT-1'} else None
    assert body['manifest']['permission_context'] == expected
    assert body['manifest']['mapping_versions'] == []
    assert body['manifest']['sources'] == []
    assert body['manifest']['approved_query'] is None
    with store.transaction() as db:
        audit = db.execute('SELECT * FROM query_audit WHERE query_id=%s', (body['query_id'],)).fetchone()
    assert audit is not None and audit['retrieval_started'] is False


def test_query_manifest_traces_seed_evidence(store: Store) -> None:
    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    started = datetime.now(timezone.utc)
    with TestClient(create_app(app)) as client:
        response = client.post('/queries', json={
            'approved_query': 'v0.1.0/overdue_balances',
            'permission': {'role': 'owner', 'client_id': 'CLIENT-1'}, 'as_of': '2026-09-01'})
    assert response.status_code == 200
    body = response.json()
    proof = body['manifest']
    assert proof['snapshot_id'] == body['snapshot_id']
    assert started <= datetime.fromisoformat(proof['retrieved_at']) <= datetime.now(timezone.utc)
    assert proof['ontology_version'] == 'v0.1.0'
    assert proof['approved_query'] == 'v0.1.0/overdue_balances'
    assert proof['metric_definition_ref'] == 'ontology/v0.1.0.yaml#metrics.overdue_balances'
    assert proof['permission_context'] == {'role': 'owner', 'client_id': 'CLIENT-1'}
    assert {'billing:invoices:I001', 'billing:payments:P001', 'billing:payments:P002',
            'crm:customers:C001'} <= set(proof['source_ids'])
    assert 'crm:customers:C002' not in proof['source_ids']
    assert 'billing:invoices:I002' not in proof['source_ids']
    snapshots = {str(s.snapshot_id) for s in store.snapshots()}
    assert proof['sources']
    assert all(s['snapshot_id'] in snapshots and s['mapping_version'] == 'v0.1.0'
               for s in proof['sources'])
    assert {'billing:v0.1.0', 'crm:v0.1.0'} <= set(proof['mapping_versions'])
    expected = json.loads(Path(__file__).with_name('query_expected.json').read_text())
    assert body['rows'][0]['observed_amount'] == expected['overdue_balances']['CLIENT-1']


def test_conflicting_support_date_is_visible_with_unresolved_ticket(store: Store) -> None:
    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    with TestClient(create_app(app)) as client:
        response = client.post('/queries', json={
            'approved_query': 'v0.1.0/overdue_balances',
            'permission': {'role': 'owner', 'client_id': 'CLIENT-1'}, 'as_of': '2026-09-01'})
        other = client.post('/queries', json={
            'approved_query': 'v0.1.0/overdue_balances',
            'permission': {'role': 'analyst', 'client_id': 'CLIENT-2'}, 'as_of': '2026-09-01'})
    proof = response.json()['manifest']
    assert any('2026-08-15' in c and '2026-08-20' in c and 'T001' in c and 'I001' in c
               for c in proof['unresolved_conflicts'])
    ticket = next(s for s in proof['sources'] if s['canonical_id'] == 'T001')
    assert ticket['ticket']['status'] == 'open'
    assert ticket['source_id'] == 'support:tickets:T001'
    assert 'support:tickets:T001' in proof['source_ids']
    assert 'support:tickets:T001' not in other.json()['manifest']['source_ids']
    assert not any('I001' in c for c in other.json()['manifest']['unresolved_conflicts'])


@pytest.fixture
def result(store: Store) -> QueryResult:
    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    create_app(app)
    return QueryService(store.schema).execute(QueryRequest(
        approved_query='v0.1.0/overdue_balances',
        permission=PermissionContext(role='owner', client_id='CLIENT-1'), as_of=date(2026, 9, 1)))


def test_markdown_briefing_has_values_inline_sources_and_complete_proof(result: QueryResult) -> None:
    from services.artifacts import generate
    from services.proof import ProofManifest

    path = Path(__file__).resolve().parents[1] / 'output-workspace' / f'{result.query_id}.md'
    assert generate(result, path) == path
    text = path.read_text()
    expected = json.loads(Path(__file__).with_name('query_expected.json').read_text())
    assert f"C001: observed {expected['overdue_balances']['CLIENT-1']}" in text
    assert '[^billing:invoices:I001]' in text
    assert 'T001' in text and 'open' in text
    assert '2026-08-15' in text and '2026-08-20' in text
    payload = text.split('```json\n')[1].split('\n```')[0]
    proof = ProofManifest.model_validate_json(payload)
    assert proof.snapshot_id == result.snapshot_id
    assert proof.source_ids == result.manifest.source_ids
    print('BRIEFING_SAMPLE=' + text)


def test_spreadsheet_cells_sources_and_schema_validated_manifest(result: QueryResult) -> None:
    from openpyxl import load_workbook
    from pydantic import ValidationError
    from services.artifacts import generate, OUTPUT
    from services.proof import ProofManifest

    path = generate(result, OUTPUT / f'{result.query_id}.xlsx')
    book = load_workbook(path)
    assert book['Results']['A2'].value == 'C001'
    expected = json.loads(Path(__file__).with_name('query_expected.json').read_text())
    assert book['Results']['C2'].value == expected['overdue_balances']['CLIENT-1']
    assert book['Results']['C2'].data_type == 's'
    assert book['Tickets']['A2'].value == 'T001'
    source_rows = list(book['Sources'].values)
    assert any(row[0] == 'billing:payments:P002' for row in source_rows)
    assert any('completeness' in str(cell) for row in source_rows for cell in row)
    assert any(str(result.manifest.sources[0].retrieved_at.isoformat()) in str(cell)
               for row in source_rows for cell in row)
    payload = json.loads(''.join(str(row[0]) for row in book['Manifest'].values))
    proof = ProofManifest.model_validate(payload)
    assert proof.snapshot_id == result.snapshot_id
    assert proof.unresolved_conflicts
    # Required fields must not acquire defaults that silently hide omissions.
    for field in ('source_ids', 'retrieved_at', 'snapshot_id', 'ontology_version', 'mapping_versions',
                  'approved_query', 'metric_definition_ref', 'permission_context', 'unresolved_conflicts'):
        with pytest.raises(ValidationError):
            ProofManifest.model_validate({key: value for key, value in payload.items() if key != field})
    book.close()


def test_deck_contains_values_tickets_and_full_manifest_in_each_notes(result: QueryResult) -> None:
    from pptx import Presentation
    from services.artifacts import generate, OUTPUT
    from services.proof import ProofManifest

    from pptx.shapes.autoshape import Shape
    deck = Presentation(str(generate(result, OUTPUT / f'{result.query_id}.pptx')))
    text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if isinstance(shape, Shape))
    assert 'C001: observed 700.00' in text  # query_expected.json seed derivation
    assert 'T001' in text and 'open' in text
    assert 'Authority unresolved' in text
    for slide in deck.slides:
        notes = slide.notes_slide.notes_text_frame
        assert notes is not None
        proof = ProofManifest.model_validate_json(notes.text)
        assert proof.snapshot_id == result.snapshot_id
        assert 'billing:payments:P002' in proof.source_ids
        assert proof.unresolved_conflicts


@pytest.mark.parametrize('extension', ['md', 'xlsx', 'pptx'])
def test_payment_absence_is_not_nonpayment(store: Store, extension: str) -> None:
    from services.artifacts import generate, OUTPUT
    from test_query_isolation import add_observation
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape

    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    add_observation(store, 'Invoice', 'I001', 'IABS', 'billing:invoices:IABS', {'amount': '25.00'})
    create_app(app)
    answer = QueryService(store.schema).execute(QueryRequest(approved_query='v0.1.0/overdue_balances',
        permission=PermissionContext(role='owner', client_id='CLIENT-1'), as_of=date(2026, 9, 1)))
    # Seed balance 700.00 plus this independently specified 25.00 invoice.
    assert str(answer.rows[0].observed_amount) == '725.00'
    assert answer.rows[0].amount is None
    assert answer.manifest.completeness_check == 'not_established'
    assert any('IABS' in c and 'no payment record found' in c for c in answer.manifest.caveats)
    assert not any('I001: no payment record found' in c for c in answer.manifest.caveats)
    path = generate(answer, OUTPUT / f'{answer.query_id}.{extension}')
    if extension == 'md':
        text = path.read_text()
    elif extension == 'xlsx':
        book = load_workbook(path)
        text = '\n'.join(str(cell) for sheet in book for row in sheet.values for cell in row)
        book.close()
    else:
        deck = Presentation(str(path))
        text = ' '.join(shape.text.replace('\n', ' ') for slide in deck.slides
                        for shape in slide.shapes if isinstance(shape, Shape))
    assert 'no payment record found' in text
    assert 'not proof of nonpayment' in text
    assert 'Completeness check: not_established' in text


@pytest.mark.parametrize('extension', ['md', 'xlsx', 'pptx'])
def test_artifact_boundary_rejects_escape_links_and_overwrite(
    result: QueryResult, tmp_path: Path, extension: str,
) -> None:
    from services.artifacts import generate, OUTPUT
    from test_ingestion import source_bytes

    before = source_bytes()
    leaf = f'{result.query_id}.{extension}'
    for outside in [tmp_path / leaf, OUTPUT.parent / leaf,
                    OUTPUT / '..' / leaf, OUTPUT.parent / 'output-workspace-evil' / leaf]:
        with pytest.raises(ValueError, match='output-workspace'):
            generate(result, outside)
        assert not outside.exists()
    OUTPUT.mkdir(exist_ok=True)
    link = OUTPUT / f'link-{result.query_id}'
    link.symlink_to(tmp_path, target_is_directory=True)
    try:
        with pytest.raises(OSError):
            generate(result, link / leaf)
        assert not (tmp_path / leaf).exists()
    finally:
        link.unlink()
    target = generate(result, OUTPUT / leaf)
    original = target.read_bytes()
    with pytest.raises(FileExistsError):
        generate(result, target)
    assert target.read_bytes() == original
    assert source_bytes() == before


def test_identity_abstention_surfaces_conflict_without_other_client_ids(store: Store) -> None:
    from test_query_isolation import add_observation

    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    add_observation(store, 'Payment', 'P001', 'PX', None, {'invoice_id': 'OTHER-CLIENT-INVOICE'})
    create_app(app)
    answer = QueryService(store.schema).execute(QueryRequest(approved_query='v0.1.0/overdue_balances',
        permission=PermissionContext(role='owner', client_id='CLIENT-1'), as_of=date(2026, 9, 1)))
    assert answer.rows == []
    assert answer.manifest.source_ids == []
    assert 'Financial identity ambiguity; metric retrieval abstained.' in answer.manifest.unresolved_conflicts
    assert 'OTHER-CLIENT-INVOICE' not in answer.model_dump_json()


@pytest.mark.parametrize('damage', ['snapshot', 'sources', 'conflicts', 'permission', 'money'])
def test_artifact_rejects_inconsistent_or_unvalidated_result(result: QueryResult, damage: str) -> None:
    from uuid import uuid4
    from pydantic import ValidationError
    from services.artifacts import generate, OUTPUT

    broken = result.model_copy(deep=True)
    if damage == 'snapshot':
        broken = broken.model_copy(update={'manifest': broken.manifest.model_copy(update={'snapshot_id': uuid4()})})
    elif damage == 'sources':
        broken.manifest.source_ids.clear()
    elif damage == 'conflicts':
        broken.manifest.unresolved_conflicts.clear()
    elif damage == 'permission':
        broken.rows[0] = broken.rows[0].model_copy(update={'client_id': 'CLIENT-2'})
    else:
        # model_copy deliberately bypasses Pydantic; generation must revalidate.
        broken.rows[0] = broken.rows[0].model_copy(update={'observed_amount': 0.1})
    path = OUTPUT / f'bad-{result.query_id}.md'
    with pytest.raises((ValueError, ValidationError)):
        generate(broken, path)
    assert not path.exists()


def test_conflict_evidence_includes_invoice_even_before_it_is_overdue(store: Store) -> None:
    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    create_app(app)
    answer = QueryService(store.schema).execute(QueryRequest(approved_query='v0.1.0/overdue_balances',
        permission=PermissionContext(role='owner', client_id='CLIENT-1'), as_of=date(2026, 8, 15)))
    assert str(answer.rows[0].observed_amount) == '0.00'
    assert any('I001' in conflict for conflict in answer.manifest.unresolved_conflicts)
    assert 'billing:invoices:I001' in answer.manifest.source_ids


def test_all_response_outcomes_have_schema_valid_proof(store: Store) -> None:
    from services.proof import ProofManifest
    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    base = {'approved_query': 'v0.1.0/cash_received',
            'permission': {'role': 'owner', 'client_id': 'CLIENT-1'}, 'as_of': '2026-09-01'}
    with TestClient(create_app(app)) as client:
        attempts: list[tuple[dict[str, object], int]] = [({'approved_query': 'v0.1.0/revenue'}, 200),
                                ({'approved_query': 'v0.1.0/pipeline'}, 200),
                                ({'approved_query': 'arbitrary SQL'}, 422),
                                ({'permission': {'role': 'guest', 'client_id': 'CLIENT-1'}}, 403),
                                ({'as_of': 'invalid'}, 422), ({'period_start': '2026-09-02'}, 422)]
        for changes, status in attempts:
            response = client.post('/queries', json={**base, **changes})
            assert response.status_code == status
            proof = ProofManifest.model_validate(response.json()['manifest'])
            assert str(proof.snapshot_id) == response.json()['snapshot_id']
            assert proof.sources == []
        with store.transaction() as db:
            db.execute('REVOKE SELECT ON canonical_records FROM lab_query')
        failure = client.post('/queries', json=base)
        assert failure.status_code == 503
        proof = ProofManifest.model_validate(failure.json()['manifest'])
        assert proof.sources == [] and proof.source_ids == []


def test_exact_large_money_and_untrusted_subject_survive_artifacts(store: Store) -> None:
    from services.artifacts import generate, OUTPUT
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    from test_query_isolation import add_observation

    app = approved_app(store)
    for family in FAMILIES:
        app.ingest(family, 'v0.1.0')
    add_observation(store, 'Payment', 'P001', 'P001', None, {'amount': '9007199254740993.01'})
    add_observation(store, 'Payment', 'P002', 'P002', None, {'amount': '-0.02'})
    add_observation(store, 'Ticket', 'T001', 'T001', None, {'subject': '=1+1 <script>alert(1)</script>'})
    create_app(app)
    answer = QueryService(store.schema).execute(QueryRequest(approved_query='v0.1.0/cash_received',
        permission=PermissionContext(role='owner', client_id='CLIENT-1'), as_of=date(2026, 9, 1)))
    # Independent decimal fixture: 9007199254740993.01 less 0.02.
    expected = '9007199254740992.99'
    book = load_workbook(generate(answer, OUTPUT / f'{answer.query_id}.xlsx'))
    assert book['Results']['C2'].value == expected
    assert book['Tickets']['C2'].value == '=1+1 <script>alert(1)</script>'
    assert book['Tickets']['C2'].data_type == 's'
    book.close()
    md = generate(answer, OUTPUT / f'{answer.query_id}.md').read_text()
    assert expected in md
    assert '<script>' not in md.split('```json')[0]
    deck = Presentation(str(generate(answer, OUTPUT / f'{answer.query_id}.pptx')))
    assert any(expected in shape.text for slide in deck.slides for shape in slide.shapes
               if isinstance(shape, Shape))
